"""Build data/regional_curves.json from the TSA3 regional-curve workbooks.

Reads the three Excel files (North Shore, Cloquet/St. Louis, Eastern MN), extracts the survey
sites and the prediction equations exactly as the spreadsheets use them, and also refits every
power-law relation from the raw rows (ordinary least squares on ln(y) vs ln(x), which is what
Excel's "power" trendline does) so the site can show whether the published coefficients still
match the data.

Usage:
  python tools/build_regional_curves.py [--src "S:/TECH/20_North_Shore_Geomorph/02-Regional_Curves"] [--out data/regional_curves.json]
"""
import argparse, json, math, os, re, sys, datetime
import openpyxl

def power_fit(xs, ys):
    pts = [(x, y) for x, y in zip(xs, ys) if x and y and x > 0 and y > 0]
    n = len(pts)
    if n < 3:
        return None
    lx = [math.log(x) for x, _ in pts]; ly = [math.log(y) for _, y in pts]
    mx = sum(lx) / n; my = sum(ly) / n
    sxx = sum((a - mx) ** 2 for a in lx); sxy = sum((a - mx) * (b - my) for a, b in zip(lx, ly))
    b = sxy / sxx; a = math.exp(my - b * mx)
    ss_res = sum((v - (math.log(a) + b * u)) ** 2 for u, v in zip(lx, ly))
    ss_tot = sum((v - my) ** 2 for v in ly)
    r2 = 1 - ss_res / ss_tot if ss_tot else None
    return {"a": round(a, 5), "b": round(b, 5), "n": n, "r2": round(r2, 4) if r2 is not None else None,
            "da_min": round(min(x for x, _ in pts), 3), "da_max": round(max(x for x, _ in pts), 3)}

def read_sites(ws, first_row=3, cols=None):
    """Rows with Stream Name in B; columns: B name, C type, D slope, E DA, F width, G depth, H W/D, I area, J desc."""
    out = []
    for r in range(first_row, ws.max_row + 1):
        name = ws.cell(r, 2).value
        if not name or not isinstance(name, str):
            continue
        da = ws.cell(r, 5).value
        if not isinstance(da, (int, float)):
            continue
        def num(c):
            v = ws.cell(r, c).value
            return round(float(v), 4) if isinstance(v, (int, float)) else None
        out.append({"stream": name.strip(), "type": (ws.cell(r, 3).value or "").strip() if isinstance(ws.cell(r, 3).value, str) else ws.cell(r, 3).value,
                    "slope": num(4), "da": num(5), "width": num(6), "depth": num(7), "area": num(9),
                    "site": ws.cell(r, 10).value})
    return out

def coef_from_formula(f):
    """Parse '=8.8667*(D1^0.5468)' → (8.8667, 0.5468)."""
    m = re.search(r"=\s*([\d.]+)\s*\*\s*\(\$?D\$?1\^([\d.]+)\)", f or "")
    return (float(m.group(1)), float(m.group(2))) if m else None

def north_shore(path):
    wbv = openpyxl.load_workbook(path, data_only=True)
    wbf = openpyxl.load_workbook(path, data_only=False)
    eq = wbf["Prediction Equations"]
    pub = {}
    for label, col in (("B", 4), ("C", 9), ("E", 14)):
        pub[label] = {"area": coef_from_formula(eq.cell(5, col).value), "width": coef_from_formula(eq.cell(7, col).value),
                      "depth": coef_from_formula(eq.cell(9, col).value), "discharge": coef_from_formula(eq.cell(13, col).value)}
    sites = {t: read_sites(wbv[f'"{t}" Channel Dimensions']) for t in ("B", "C", "E")}
    all_sites = read_sites(wbv["ALL Channels - Area"])
    q_ws = wbv["Discharge"]
    q_sites = []
    for r in range(4, 19):
        n, da, q = q_ws.cell(r, 2).value, q_ws.cell(r, 3).value, q_ws.cell(r, 4).value
        if isinstance(da, (int, float)) and isinstance(q, (int, float)):
            q_sites.append({"stream": n, "da": da, "q": q})
    fits = {}
    for t, rows in sites.items():
        fits[t] = {"area": power_fit([s["da"] for s in rows], [s["area"] for s in rows]),
                   "width": power_fit([s["da"] for s in rows], [s["width"] for s in rows]),
                   "depth": power_fit([s["da"] for s in rows], [s["depth"] for s in rows])}
    fits["ALL"] = {"area": power_fit([s["da"] for s in all_sites], [s["area"] for s in all_sites])}
    fits["discharge"] = power_fit([s["da"] for s in q_sites], [s["q"] for s in q_sites])
    return {
        "id": "north_shore", "name": "North Shore Regional Curve",
        "region": "Lake Superior North Shore tributaries: Duluth northeast through Cook County (HUC8 04010101, 04010102)",
        "source_file": os.path.basename(path), "source_modified": datetime.datetime.fromtimestamp(os.path.getmtime(path)).date().isoformat(),
        "reliability": "Primary curve for TSA3. Best-supported dataset; B, C and E channel types fit separately.",
        "method": "Power-law fits of bankfull area, width and mean depth vs drainage area per Rosgen stream type; bankfull discharge from gage-site surveys (single fit for all types); velocity = Q / area.",
        "types": {
            t: {"label": {"B": "B channels (steep, confined)", "C": "C channels (riffle-pool, moderate W/D)", "E": "E channels (low W/D, meandering)"}[t],
                "equations": {"area": {"form": "power", "a": pub[t]["area"][0], "b": pub[t]["area"][1], "unit": "ft²"},
                              "width": {"form": "power", "a": pub[t]["width"][0], "b": pub[t]["width"][1], "unit": "ft"},
                              "depth": {"form": "power", "a": pub[t]["depth"][0], "b": pub[t]["depth"][1], "unit": "ft"},
                              "discharge": {"form": "power", "a": pub[t]["discharge"][0], "b": pub[t]["discharge"][1], "unit": "cfs"}},
                "refit": fits[t], "sites": sites[t]}
            for t in ("B", "C", "E")},
        "all_sites": all_sites, "all_area_refit": fits["ALL"]["area"],
        "discharge_sites": q_sites, "discharge_refit": fits["discharge"],
    }

def cloquet(path):
    wbv = openpyxl.load_workbook(path, data_only=True)
    wbf = openpyxl.load_workbook(path, data_only=False)
    eq = wbf["Prediction Equations"]
    area = coef_from_formula(eq.cell(5, 4).value)
    wd = {"B": eq.cell(11, 4).value, "C": eq.cell(11, 9).value, "E": eq.cell(11, 14).value}
    vel = float(re.search(r"\*\s*([\d.]+)", eq.cell(13, 4).value).group(1))
    sites = read_sites(wbv["ALL Channels Area"])
    return {
        "id": "cloquet_st_louis", "name": "Cloquet / St. Louis River Regional Curve",
        "region": "Cloquet River and St. Louis River watersheds (HUC8 04010201, 04010202) and the Nemadji headwaters in Carlton County",
        "source_file": os.path.basename(path), "source_modified": datetime.datetime.fromtimestamp(os.path.getmtime(path)).date().isoformat(),
        "reliability": "Single bankfull-area regression across all channel types (15 sites, several area-only). Width and depth are NOT regressed: they come from an assumed W/D ratio per stream type (B 18, C 20, E 12). Discharge assumes a bankfull velocity of 3.5 ft/s.",
        "method": "Area = a·DA^b for all types; width = sqrt(area × W/D); depth = width / W/D; Q = area × 3.5 ft/s.",
        "types": {t: {"label": {"B": "B channels (W/D 18 assumed)", "C": "C channels (W/D 20 assumed)", "E": "E channels (W/D 12 assumed)"}[t],
                      "equations": {"area": {"form": "power", "a": area[0], "b": area[1], "unit": "ft²"},
                                    "wd_ratio": wd[t], "velocity_fps": vel}} for t in ("B", "C", "E")},
        "all_sites": sites, "all_area_refit": power_fit([s["da"] for s in sites], [s["area"] for s in sites]),
    }

def eastern(path):
    wbf = openpyxl.load_workbook(path, data_only=False)
    wbv = openpyxl.load_workbook(path, data_only=True)
    eq = wbf["Equations"]
    def poly_from(f):
        # coefficients of a polynomial in D1, highest power first; handles '=0.0773*(D1^3)-1.1993*(D1^2)+6.0889*(D1)+0.5281'
        f = f.replace(" ", "").lstrip("=")
        terms = re.findall(r"([+-]?[\d.eE-]+)\*?\(?D1\^?(\d?)\)?|([+-]?[\d.]+)$", f)
        coefs = {}
        for c, p, const in terms:
            if const:
                coefs[0] = float(const)
            else:
                coefs[int(p) if p else 1] = float(c)
        deg = max(coefs)
        return [coefs.get(d, 0.0) for d in range(deg, -1, -1)]
    area_lt5 = poly_from(eq.cell(3, 5).value)
    area_ge5 = coef_from_formula(eq.cell(5, 5).value)
    width = coef_from_formula(eq.cell(7, 5).value)
    depth = coef_from_formula(eq.cell(9, 5).value)
    q_poly = poly_from(eq.cell(11, 5).value)
    def table(ws, c1=1, c2=2):
        return [{"da": ws.cell(r, c1).value, "v": ws.cell(r, c2).value} for r in range(2, ws.max_row + 1) if isinstance(ws.cell(r, c1).value, (int, float))]
    return {
        "id": "eastern_mn", "name": "Eastern MN Regional Curve",
        "region": "Kanabec, Mille Lacs and Pine counties (Snake, Rum, Kettle and upper St. Croix basins)",
        "source_file": os.path.basename(path), "source_modified": datetime.datetime.fromtimestamp(os.path.getmtime(path)).date().isoformat(),
        "reliability": "Least reliable of the three: equations are fit to smoothed curve-read values, not to individual surveyed sites, and there is no stream-type split. The North Shore C and E curves are often closer for these counties; compare both.",
        "method": "Bankfull area: cubic polynomial below 5 mi², power law at or above 5 mi². Width and depth: power laws. Discharge: 4th-order polynomial. All fit to digitized curve tables.",
        "equations": {"area_lt5": {"form": "poly", "coefs": area_lt5, "unit": "ft²", "valid": "DA < 5 mi²"},
                      "area_ge5": {"form": "power", "a": area_ge5[0], "b": area_ge5[1], "unit": "ft²", "valid": "DA ≥ 5 mi²"},
                      "width": {"form": "power", "a": width[0], "b": width[1], "unit": "ft"},
                      "depth": {"form": "power", "a": depth[0], "b": depth[1], "unit": "ft"},
                      "discharge": {"form": "poly", "coefs": q_poly, "unit": "cfs"}},
        "tables": {"area": table(wbv["BKF Area"], 2, 3), "discharge": table(wbv["Discharge"]), "width": table(wbv["Width"]), "depth": table(wbv["Mean Depth"])},
    }

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default="S:/TECH/20_North_Shore_Geomorph/02-Regional_Curves")
    ap.add_argument("--out", default="data/regional_curves.json")
    a = ap.parse_args()
    ns = north_shore(os.path.join(a.src, "North Shore Regional Curve_updated 2026_01.xlsx"))
    cl = cloquet(os.path.join(a.src, "Cloquet St Louis Regional Curve_2019_11.xlsx"))
    em = eastern(os.path.join(a.src, "EasternMN_Regional_Curve.xlsx"))
    out = {"generated": datetime.datetime.now().isoformat(timespec="seconds"), "generator": "tools/build_regional_curves.py",
           "units": {"da": "mi²", "area": "ft²", "width": "ft", "depth": "ft", "discharge": "cfs", "velocity": "ft/s"},
           "curves": [ns, cl, em]}
    os.makedirs(os.path.dirname(a.out), exist_ok=True)
    with open(a.out, "w", encoding="utf-8") as f:
        json.dump(out, f, indent=1)
    # Report: published vs refit
    print(f"wrote {a.out}")
    print("\nNorth Shore: published coefficients vs refit from current survey rows")
    for t in ("B", "C", "E"):
        for k in ("area", "width", "depth"):
            p = ns["types"][t]["equations"][k]; r = ns["types"][t]["refit"][k]
            print(f"  {t} {k:6s} published a={p['a']:<8} b={p['b']:<7}  refit a={r['a']:<9} b={r['b']:<8} n={r['n']} R²={r['r2']}  DA {r['da_min']}–{r['da_max']}")
    p = ns["types"]["C"]["equations"]["discharge"]; r = ns["discharge_refit"]
    print(f"  Q (all)  published a={p['a']} b={p['b']}  refit a={r['a']} b={r['b']} n={r['n']} R²={r['r2']}")
    r = ns["all_area_refit"]; print(f"  ALL area refit a={r['a']} b={r['b']} n={r['n']} R²={r['r2']}")
    print("\nCloquet/St. Louis: published area", cl["types"]["C"]["equations"]["area"], "refit", cl["all_area_refit"])
    print("Eastern MN equations:", json.dumps(em["equations"]))

if __name__ == "__main__":
    main()
